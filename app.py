from datetime import date, datetime
from pathlib import Path
import streamlit as st
import pandas as pd
import openpyxl

BASE_DIR = Path(__file__).resolve().parent
EXCEL_FILE = BASE_DIR / "Operação.xlsx"
LOGO_FILE = BASE_DIR / "assets" / "logo.png"
HERO_FILE = BASE_DIR / "assets" / "FotoUD.jpg"
REALIZADORES_FILE = BASE_DIR / "assets" / "realizadores.png"
HISTORY_FILE = BASE_DIR / "historia.json"
HISTORY_PHOTOS = BASE_DIR / "assets" / "historico"

st.set_page_config(
    page_title="Unidade de Biocombustíveis de Itaipu",
    page_icon="🌱",
    layout="wide",
    initial_sidebar_state="expanded",
)

st.markdown("""
<style>
.block-container {padding-top:1.3rem;padding-bottom:2rem;max-width:1500px;}
section[data-testid="stSidebar"] {border-right:1px solid #E1E9E4;}
.hero {padding:1.8rem 2rem;border-radius:16px;color:#fff;
background:linear-gradient(120deg,#174B37 0%,#1F6B4A 62%,#4A8A69 100%);margin-bottom:1.2rem;}
.site-hero {min-height:440px;display:flex;align-items:center;padding:clamp(1.25rem,4vw,3.5rem);border-radius:18px;color:white;background-size:cover;background-position:center;margin:0 0 1.5rem 0;box-sizing:border-box;}
.site-hero-content {display:flex;align-items:center;gap:clamp(1.25rem,3vw,2.5rem);max-width:1120px;}
.site-hero-logo {background:rgba(255,255,255,.97);padding:1.1rem;border-radius:14px;flex:0 0 300px;box-shadow:0 8px 30px rgba(0,0,0,.16);box-sizing:border-box;}
.site-hero-logo img {display:block;width:100%;height:auto;object-fit:contain;}
.site-hero-copy {max-width:720px;text-shadow:0 1px 8px rgba(0,0,0,.18);}
.site-hero h1 {font-size:clamp(2rem,4vw,3.7rem);line-height:1.08;max-width:850px;margin:.5rem 0 1rem 0;}
.site-hero p {font-size:1.1rem;max-width:720px;}
.site-kicker {font-size:.8rem;letter-spacing:.15em;font-weight:800;text-transform:uppercase;}
@media (max-width:760px){.site-hero{min-height:unset;padding:1.25rem}.site-hero-content{flex-direction:column;align-items:flex-start;gap:1.25rem}.site-hero-logo{flex:0 1 auto;width:min(100%,300px);padding:.75rem}.site-hero h1{font-size:2.15rem}.site-hero p{font-size:1rem}}
.site-footer {margin-top:3rem;padding:1.5rem 0 .5rem;border-top:1px solid #dce5df;text-align:center;}
.footer-logos {display:flex;align-items:center;justify-content:center;gap:3rem;flex-wrap:wrap;margin:1.25rem auto .5rem auto;}
.footer-logos img {max-width:230px;max-height:76px;width:auto;height:auto;object-fit:contain;}
@media (max-width:640px){.footer-logos{gap:1.5rem}.footer-logos img{max-width:175px;max-height:60px}}
.history-card {padding:1.1rem 1.25rem;border-left:4px solid #26734d;background:#f5f8f6;border-radius:0 12px 12px 0;margin-bottom:1rem;}
.hero h1 {font-size:clamp(1.8rem,3vw,2.65rem);line-height:1.12;margin:0 0 .65rem 0;}
.hero p {font-size:1.03rem;max-width:900px;margin:0;opacity:.96;}
.eyebrow {font-size:.76rem;letter-spacing:.11em;text-transform:uppercase;font-weight:700;color:#D4E9DD;}
.muted {color:#617168;}
div[data-testid="stMetric"] {background:#F5F8F6;border:1px solid #E2EAE5;padding:13px 15px;border-radius:12px;}
</style>
""", unsafe_allow_html=True)

VIDEO_LINKS = {
    "Institucional": "https://youtu.be/qfgAGSnDA4U",
    "Segurança": "https://youtu.be/RVSyaAXuDuw",
    "Diário de obra": "https://youtu.be/8wWsIYfZUf8",
}

def brand(compact=False):
    a, b = st.columns([1, 4.5], vertical_alignment="center")
    with a:
        if LOGO_FILE.exists():
            st.image(str(LOGO_FILE), use_container_width=True)
        else:
            st.markdown("**UNIDADE DE BIOCOMBUSTÍVEIS**")
    with b:
        st.markdown("### Unidade de Biocombustíveis de Itaipu")
        st.caption("Biogás • Biometano • Pesquisa aplicada • Transição energética")

def read_excel_data():
    """Leitura estritamente somente leitura do arquivo Excel, sem automação do Excel/COM."""
    if not EXCEL_FILE.exists():
        raise FileNotFoundError(f"Arquivo não encontrado: {EXCEL_FILE.name}")
    wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True, read_only=True)
    try:
        feeding = []
        if "Operação_CSTR" in wb.sheetnames:
            ws = wb["Operação_CSTR"]
            for row in ws.iter_rows(min_row=2, max_row=min(ws.max_row, 500), min_col=1, max_col=99, values_only=True):
                raw_date = row[0]
                if not isinstance(raw_date, (datetime, date)):
                    continue
                d = raw_date.date() if isinstance(raw_date, datetime) else raw_date
                def num(idx):
                    v = row[idx] if idx < len(row) else None
                    return float(v) if isinstance(v, (int, float)) else 0.0
                restaurant = num(20)  # U: massa do resíduo de restaurante
                additional = 0.0
                # Y, AC, AG...CS: massa dos substratos adicionais.
                for idx in range(24, 97, 4):
                    additional += num(idx)
                feeding.append({
                    "Data": pd.Timestamp(d),
                    "Volume do reator (m³)": num(4),
                    "Biogás estimado (m³)": num(5),
                    "Diluição/recirculação (L)": num(17),
                    "Resíduo de restaurante (kg)": num(18),
                    "Massa do restaurante registrada (kg)": restaurant,
                    "Outros substratos (kg)": additional,
                    "Alimentação total calculada (kg)": restaurant + additional,
                })
        gas = []
        if "Composição_Biogás" in wb.sheetnames:
            ws = wb["Composição_Biogás"]
            # Cabeçalho na linha 7; registros esperados da linha 8 até 205.
            for row in ws.iter_rows(min_row=8, max_row=min(ws.max_row, 205), min_col=1, max_col=13, values_only=True):
                raw_date = row[0]
                if not isinstance(raw_date, (datetime, date)):
                    continue
                d = raw_date.date() if isinstance(raw_date, datetime) else raw_date
                def val(idx):
                    v = row[idx] if idx < len(row) else None
                    if isinstance(v, (int, float)):
                        return float(v)
                    return None
                gas.append({
                    "Data": pd.Timestamp(d),
                    "Horário": str(row[1] or ""),
                    "Ponto de medição": str(row[2] or "").strip(),
                    "CH₄ (%)": val(3),
                    "CO₂ (%)": val(4),
                    "O₂ (%)": val(5),
                    "NH₃ (ppm)": val(6),
                    "CO (ppm)": val(7),
                    "H₂S (ppm)": val(8),
                    "Totalizador (m³)": val(10),
                    "Produção indicada (m³)": val(11),
                    "Observação": str(row[12] or ""),
                })
        fdf = pd.DataFrame(feeding)
        gdf = pd.DataFrame(gas)
        if not gdf.empty:
            gdf = gdf.sort_values(["Data", "Horário"]).reset_index(drop=True)
            gdf["Volume entre leituras (m³)"] = gdf["Totalizador (m³)"].diff()
            gdf.loc[gdf["Volume entre leituras (m³)"] < 0, "Volume entre leituras (m³)"] = None
        return fdf, gdf
    finally:
        wb.close()

def show_footer():
    st.markdown("""
    <div class="site-footer">
      <div style="color:#587064;font-size:.85rem">Unidade de Biocombustíveis de Itaipu · Biogás, biometano e inovação</div>
      <div style="color:#587064;font-size:.75rem;margin-top:1rem;text-transform:uppercase;letter-spacing:.08em">Realização</div>
      <div class="footer-logos">
        <img src="https://cibiogas.org/wp-content/uploads/2022/06/LOGO.svg" alt="CIBiogás" loading="lazy">
        <img src="https://www.itaipu.gov.br/wp-content/uploads/2024/03/Logo_Itaipu_Preferencial.svg-1024x724-1.png" alt="Itaipu Binacional" loading="lazy">
      </div>
    </div>
    """, unsafe_allow_html=True)

def show_home():
    # Testeira institucional: fotografia real da unidade, logo oficial e título integrados.
    import base64
    if HERO_FILE.exists():
        encoded_photo = base64.b64encode(HERO_FILE.read_bytes()).decode("ascii")
        background = f"linear-gradient(90deg, rgba(9,42,29,.86) 0%, rgba(12,57,39,.68) 58%, rgba(12,57,39,.24) 100%), url(data:image/jpeg;base64,{encoded_photo})"
    else:
        background = "linear-gradient(120deg,#174B37,#4A8A69)"
    if LOGO_FILE.exists():
        encoded_logo = base64.b64encode(LOGO_FILE.read_bytes()).decode("ascii")
        logo_html = f'<div class="site-hero-logo"><img src="data:image/png;base64,{encoded_logo}" alt="Logo da Unidade de Demonstração de Biocombustíveis de Itaipu"></div>'
    else:
        logo_html = '<div class="site-hero-logo"><strong>UNIDADE DE BIOCOMBUSTÍVEIS · ITAIPU</strong></div>'
    st.markdown(f'''<section class="site-hero" style="background-image:{background}">
      <div class="site-hero-content">{logo_html}
        <div class="site-hero-copy">
          <div class="site-kicker">ITAIPU · ENERGIA · INOVAÇÃO</div>
          <h1>Unidade de Biocombustíveis</h1>
          <p>Transformando resíduos em energia e conhecimento. Um espaço de demonstração, pesquisa aplicada e desenvolvimento de soluções em biogás e biometano.</p>
        </div>
      </div>
    </section>''', unsafe_allow_html=True)
    st.markdown("## Conheça a unidade")
    st.write("Explore a história da unidade, conheça os conteúdos audiovisuais e acompanhe os registros operacionais disponíveis.")
    c1, c2, c3 = st.columns(3)
    with c1:
        st.markdown("#### Nossa história")
        st.write("Conheça os marcos, as melhorias e os momentos que construíram a trajetória da unidade.")
        if st.button("Explorar histórico", key="home_history", use_container_width=True):
            st.session_state["nav_page"] = "Histórico"
            st.rerun()
    with c2:
        st.markdown("#### Vídeos")
        st.write("Apresentações institucionais, segurança e registros das obras.")
        if st.button("Acessar vídeos", key="home_videos", use_container_width=True):
            st.session_state["nav_page"] = "Vídeos"
            st.rerun()
    with c3:
        st.markdown("#### Dashboard operacional")
        st.write("Consulte indicadores de alimentação, produção estimada e composição do biogás registrados na planilha.")
        if st.button("Acessar dashboard", type="primary", key="home_dashboard", use_container_width=True):
            st.session_state["nav_page"] = "Dashboard operacional"
            st.rerun()
    show_footer()

def show_history():
    brand(compact=True)
    st.title("Histórico da unidade")
    st.write("Uma linha do tempo da trajetória, das melhorias e dos momentos importantes da Unidade de Biocombustíveis de Itaipu.")
    st.info("Este histórico começa com um texto-base editável. Para registrar datas e acontecimentos com precisão, edite o arquivo historia.json e adicione fotos à pasta assets/historico.")
    import json
    try:
        entries = json.loads(HISTORY_FILE.read_text(encoding="utf-8"))
    except Exception as exc:
        st.error(f"Não foi possível ler historia.json: {exc}")
        return
    for entry in entries:
        st.markdown('<div class="history-card">', unsafe_allow_html=True)
        st.markdown(f"### {entry.get('ano','Marco da trajetória')} — {entry.get('titulo','')}")
        st.write(entry.get("texto", ""))
        photo_names = entry.get("fotos", [])
        existing = [HISTORY_PHOTOS / name for name in photo_names if (HISTORY_PHOTOS / name).exists()]
        if existing:
            cols = st.columns(min(3, len(existing)))
            for i, photo in enumerate(existing):
                with cols[i % len(cols)]:
                    st.image(str(photo), use_container_width=True, caption=photo.stem.replace("_", " "))
        st.markdown('</div>', unsafe_allow_html=True)
    st.markdown("### Galeria de fotos")
    gallery = sorted([p for p in HISTORY_PHOTOS.glob("*") if p.suffix.lower() in {".jpg", ".jpeg", ".png", ".webp"}])
    if gallery:
        cols = st.columns(3)
        for i, photo in enumerate(gallery):
            with cols[i % 3]:
                st.image(str(photo), use_container_width=True, caption=photo.stem.replace("_", " "))
    else:
        st.caption("As fotos históricas aparecerão aqui quando forem adicionadas à pasta assets/historico.")
    show_footer()

def show_videos():
    brand(compact=True)
    st.title("Biblioteca de vídeos")
    st.write("Conteúdos institucionais e operacionais da Unidade de Biocombustíveis de Itaipu.")
    labels = list(VIDEO_LINKS.keys())
    current = st.session_state.get("video_choice", "Institucional")
    idx = labels.index(current) if current in labels else 0
    chosen = st.radio("Selecione o conteúdo", labels, horizontal=True, index=idx)
    st.session_state["video_choice"] = chosen
    st.video(VIDEO_LINKS[chosen])
    notes = {
        "Institucional": "Apresentação institucional da unidade.",
        "Segurança": "Conteúdo audiovisual sobre segurança.",
        "Diário de obra": "Registro audiovisual das obras e melhorias.",
    }
    st.caption(notes[chosen])
    st.markdown(f"[Abrir vídeo diretamente no YouTube]({VIDEO_LINKS[chosen]})")

@st.cache_data(ttl=60, show_spinner=False)
def cached_excel_data(file_mtime):
    # file_mtime invalida o cache quando o arquivo Excel for atualizado.
    return read_excel_data()

def show_dashboard():
    brand(compact=True)
    st.title("Dashboard operacional")
    st.caption("Fonte: arquivo Excel local • Somente leitura • Cache de 60 segundos")
    if not EXCEL_FILE.exists():
        st.error("O arquivo Operação.xlsx não está na mesma pasta do aplicativo.")
        return
    if st.button("Atualizar dados do Excel", key="refresh_excel"):
        cached_excel_data.clear()
        st.rerun()
    try:
        fdf, gdf = cached_excel_data(EXCEL_FILE.stat().st_mtime)
    except Exception as exc:
        st.error("Não foi possível ler a planilha Excel.")
        with st.expander("Detalhes técnicos"):
            st.code(str(exc))
        return
    if fdf.empty and gdf.empty:
        st.warning("Não foram encontradas datas válidas nas abas configuradas.")
        st.caption("Abas esperadas: Operação_CSTR e Composição_Biogás.")
        return

    all_dates = []
    if not fdf.empty:
        all_dates.extend(fdf["Data"].dropna().tolist())
    if not gdf.empty:
        all_dates.extend(gdf["Data"].dropna().tolist())
    dmin = min(all_dates).date()
    dmax = max(all_dates).date()

    with st.sidebar:
        st.markdown("### Filtros")
        period = st.date_input("Período", value=(dmin, dmax), min_value=dmin, max_value=dmax, key="dashboard_period")
        locs = sorted(x for x in (gdf["Ponto de medição"].dropna().unique() if not gdf.empty else []) if x)
        selected_locs = st.multiselect("Ponto de medição", locs, default=locs, key="dashboard_locations")

    if isinstance(period, (tuple, list)) and len(period) == 2:
        start, end = pd.Timestamp(period[0]), pd.Timestamp(period[1])
    else:
        start = end = pd.Timestamp(period)
    f = fdf[(fdf["Data"] >= start) & (fdf["Data"] <= end)] if not fdf.empty else fdf
    g = gdf[(gdf["Data"] >= start) & (gdf["Data"] <= end)] if not gdf.empty else gdf
    if not g.empty and selected_locs:
        g = g[g["Ponto de medição"].isin(selected_locs)]

    def metric_sum(df, col):
        return df[col].sum() if not df.empty and col in df else None
    def fmt(x, suffix=""):
        if x is None or pd.isna(x):
            return "—"
        return f"{x:,.1f}{suffix}".replace(",", "X").replace(".", ",").replace("X", ".")
    k1, k2, k3, k4 = st.columns(4)
    k1.metric("Alimentação calculada", fmt(metric_sum(f, "Alimentação total calculada (kg)"), " kg"))
    k2.metric("Resíduo de restaurante", fmt(metric_sum(f, "Massa do restaurante registrada (kg)"), " kg"))
    k3.metric("CH₄ médio", fmt(g["CH₄ (%)"].mean() if not g.empty else None, " %"))
    k4.metric("H₂S máximo", fmt(g["H₂S (ppm)"].max() if not g.empty else None, " ppm"))

    tab1, tab2, tab3 = st.tabs(["Alimentação CSTR", "Composição do biogás", "Dados tabulares"])
    with tab1:
        st.subheader("Gráficos de alimentação")
        if f.empty:
            st.info("Sem registros no período selecionado.")
        else:
            st.bar_chart(f.set_index("Data")[["Massa do restaurante registrada (kg)", "Outros substratos (kg)"]])
    with tab2:
        st.subheader("Composição do biogás e totalizador")
        if g.empty:
            st.info("Sem medições no período/ponto selecionado.")
        else:
            cols_pct = [c for c in ["CH₄ (%)", "CO₂ (%)", "O₂ (%)"] if c in g]
            st.line_chart(g.set_index("Data")[cols_pct])
            cols_ppm = [c for c in ["H₂S (ppm)", "NH₃ (ppm)", "CO (ppm)"] if c in g]
            st.line_chart(g.set_index("Data")[cols_ppm])
            st.info("O volume entre leituras é calculado pela diferença entre valores consecutivos do totalizador. O intervalo pode atravessar dias sem medição; portanto, não deve ser interpretado automaticamente como produção de um único dia.")
            show_cols = ["Data", "Horário", "Ponto de medição", "Totalizador (m³)", "Volume entre leituras (m³)", "Observação"]
            st.dataframe(g[show_cols].sort_values("Data", ascending=False), use_container_width=True, hide_index=True)
            st.download_button("Baixar medições de biogás CSV", g.to_csv(index=False).encode("utf-8-sig"), "composicao_biogas.csv", "text/csv")
    with tab3:
        a, b = st.columns(2)
        with a:
            st.markdown("**Operação CSTR**")
            st.dataframe(f, use_container_width=True, hide_index=True)
            st.download_button("Exportar CSTR CSV", f.to_csv(index=False).encode("utf-8-sig"), "operacao_cstr.csv", "text/csv", key="exp_cstr")
        with b:
            st.markdown("**Composição do biogás**")
            st.dataframe(g, use_container_width=True, hide_index=True)
            st.download_button("Exportar biogás CSV", g.to_csv(index=False).encode("utf-8-sig"), "composicao_biogas.csv", "text/csv", key="exp_gas")

def main():
    with st.sidebar:
        st.markdown("## Portal da Unidade")
        st.radio("Navegação", ["Página inicial", "Histórico", "Vídeos", "Dashboard operacional"], key="nav_page", label_visibility="collapsed")
        st.divider()
        st.caption("Unidade de Biocombustíveis de Itaipu")
        st.caption("Portal institucional e operacional")
    page = st.session_state.get("nav_page", "Página inicial")
    if page == "Página inicial":
        show_home()
    elif page == "Histórico":
        show_history()
    elif page == "Vídeos":
        show_videos()
    else:
        show_dashboard()

if __name__ == "__main__":
    main()
